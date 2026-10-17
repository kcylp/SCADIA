# -*- coding: utf-8 -*-
"""SCADIA 两份独立 Word 手册生成器（正式排版）。
输出：D:/SCADIA-工作站迁移/SCADIA-delivery/docs/手册/*.docx
"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = 'D:/SCADIA-工作站迁移/SCADIA-delivery/docs'
IMG  = BASE + '/手册/images'
OUT  = BASE + '/手册'
os.makedirs(OUT, exist_ok=True)

CJK = '微软雅黑'
CJK_BODY = '宋体'
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT2 = RGBColor(0x2E, 0x75, 0xB6)
GREY = RGBColor(0x59, 0x59, 0x59)

# ---------------- low level helpers ----------------
def _cjk(run, font=CJK_BODY, size=None, bold=None, color=None, mono=False):
    f = 'Consolas' if mono else font
    run.font.name = f
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rF = rPr.find(qn('w:rFonts'))
    if rF is None:
        rF = OxmlElement('w:rFonts'); rPr.append(rF)
    rF.set(qn('w:ascii'), f); rF.set(qn('w:hAnsi'), f)
    rF.set(qn('w:eastAsia'), 'Consolas' if mono else font)

def add_field(paragraph, instr):
    r = paragraph.add_run()
    fc1 = OxmlElement('w:fldChar'); fc1.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    fc2 = OxmlElement('w:fldChar'); fc2.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = '（在 Word 中按 F9 或右键“更新域”刷新）'
    fc3 = OxmlElement('w:fldChar'); fc3.set(qn('w:fldCharType'), 'end')
    r._r.append(fc1); r._r.append(it); r._r.append(fc2); r._r.append(t); r._r.append(fc3)
    _cjk(r, CJK_BODY, 10.5, color=GREY)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def cell_text(cell, text, bold=False, size=10.5, color=None, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    if align is not None: p.alignment = align
    for i, line in enumerate(str(text).split('\n')):
        if i: p = cell.add_paragraph()
        r = p.add_run(line); _cjk(r, CJK_BODY, size, bold, color)
    return cell

# ---------------- document scaffolding ----------------
def build_doc(manual_title, subtitle, version, meta_rows, footer_title):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = CJK_BODY; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), CJK_BODY)
    st.paragraph_format.line_spacing = 1.4
    st.paragraph_format.space_after = Pt(4)

    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.6); sec.right_margin = Cm(2.4)

    # ---- cover ----
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cjk(p.add_run('SCADIA'), CJK, 30, True, ACCENT)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cjk(p.add_run('工业监控与八场景视频分析系统'), CJK, 16, False, ACCENT2)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    _cjk(p.add_run(manual_title), CJK, 26, True, RGBColor(0,0,0))
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cjk(p.add_run(subtitle), CJK, 12, False, GREY)
    # rule
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single'); bottom.set(qn('w:sz'), '12'); bottom.set(qn('w:space'), '1'); bottom.set(qn('w:color'), '1F4E79')
    pbdr.append(bottom); pPr.append(pbdr)
    for _ in range(6):
        doc.add_paragraph()
    t = doc.add_table(rows=0, cols=2); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for k, v in meta_rows:
        row = t.add_row().cells
        cell_text(row[0], k, True, 10.5, ACCENT); shade(row[0], 'EAF1F8')
        cell_text(row[1], v, False, 10.5)
        row[0].width = Cm(4.2); row[1].width = Cm(9.0)
    doc.add_page_break()

    # ---- 版本记录 / 阅读说明 ----
    h = doc.add_heading(level=1); _cjk(h.add_run('文档信息'), CJK, 18, True, ACCENT)
    return doc

def start_body(doc, footer_title, toc=True):
    # footer with page number
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cjk(fp.add_run(footer_title + '    第 '), CJK_BODY, 9, color=GREY)
    add_field(fp, 'PAGE')
    _cjk(fp.add_run(' 页'), CJK_BODY, 9, color=GREY)
    hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _cjk(hp.add_run(footer_title), CJK_BODY, 9, color=GREY)
    if toc:
        h = doc.add_heading(level=1); _cjk(h.add_run('目录'), CJK, 18, True, ACCENT)
        p = doc.add_paragraph(); add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
        doc.add_page_break()

# ---------------- content renderer ----------------
def h1(doc, text):
    p = doc.add_heading(level=1); _cjk(p.add_run(text), CJK, 17, True, ACCENT)
    p.paragraph_format.space_before = Pt(14); p.paragraph_format.space_after = Pt(6)
def h2(doc, text):
    p = doc.add_heading(level=2); _cjk(p.add_run(text), CJK, 14, True, ACCENT2)
def h3(doc, text):
    p = doc.add_heading(level=3); _cjk(p.add_run(text), CJK, 12.5, True, RGBColor(0x33,0x33,0x33))
def para(doc, text):
    p = doc.add_paragraph(); _cjk(p.add_run(text), CJK_BODY, 10.5)
    p.paragraph_format.first_line_indent = Cm(0)
def bullets(doc, items, numbered=False):
    for it in items:
        p = doc.add_paragraph(style='List Number' if numbered else 'List Bullet')
        _cjk(p.add_run(it), CJK_BODY, 10.5)
def note(doc, text, color=RGBColor(0xB0,0x30,0x00)):
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left'); left.set(qn('w:val'),'single'); left.set(qn('w:sz'),'18'); left.set(qn('w:space'),'6'); left.set(qn('w:color'),'B03000')
    pbdr.append(left); pPr.append(pbdr)
    p.paragraph_format.left_indent = Cm(0.3)
    _cjk(p.add_run('提示  '), CJK, 10, True, color)
    _cjk(p.add_run(text), CJK_BODY, 10)
def table(doc, rows, widths=None, header=True):
    t = doc.add_table(rows=0, cols=len(rows[0])); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cell_text(cells[ci], val, bold=(header and ri == 0), size=10,
                      color=(RGBColor(0xFF,0xFF,0xFF) if header and ri == 0 else None))
            if header and ri == 0: shade(cells[ci], '2E75B6')
            elif ri % 2 == 0: shade(cells[ci], 'F4F8FC')
            if widths: cells[ci].width = Cm(widths[ci])
    doc.add_paragraph()
def figure(doc, img_id, caption, width_cm=15.5):
    path = os.path.join(IMG, img_id + '.png')
    if not os.path.exists(path):
        note(doc, f'（截图 {img_id} 缺失）'); return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _cjk(c.add_run(caption), CJK_BODY, 9, False, GREY)
def pagebreak(doc):
    doc.add_page_break()
